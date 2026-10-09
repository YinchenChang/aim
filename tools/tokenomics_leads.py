#!/usr/bin/env python3
# BY: SCOUT@CL | MODEL: Claude Opus 5.5【自報】 | CHARTER: AIM-C1.2 | EMITS: tools/tokenomics_leads.py
"""tools/tokenomics_leads.py — read-only lead index from Andy's Tokenomics repo.

Exports WHERE to look (SRC_ID, indicator, date, source, grade, stance), never
the model's values, ranges, conclusions or Andy's own (grade A) entries.
Every lead must be re-fetched from its original source before use.

Usage:
  python tools/tokenomics_leads.py "egress" "telemetry" [--sheet SRC_Price] [--n 30]
  python tools/tokenomics_leads.py --list-sheets
"""
import argparse
import os
import subprocess
import sys

REPO = "https://github.com/YinchenChang/tokenomics"
DEST = "/home/claude/yinchenchang/tokenomics"   # read-only shallow clone
KEEP = ["SRC_ID", "指標", "口徑", "適用對象", "日期", "出處", "來源等級", "立場", "一手／二手"]


def ensure_clone() -> str:
    if not os.path.isdir(os.path.join(DEST, ".git")):
        os.makedirs(os.path.dirname(DEST), exist_ok=True)
        env = dict(os.environ, GIT_LFS_SKIP_SMUDGE="1")
        subprocess.run(["git", "clone", "-q", "--depth", "1", REPO, DEST],
                       check=True, env=env, timeout=600)
    cur = open(os.path.join(DEST, "model", "CURRENT")).read().strip()
    return os.path.join(DEST, "model", cur)


def rows(path: str, sheet_filter: str | None):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        if not name.startswith("SRC_") or name == "SRC_Index":
            continue
        if sheet_filter and name != sheet_filter:
            continue
        hdr = None
        for r in wb[name].iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in r]
            if hdr is None:
                if "SRC_ID" in cells:
                    hdr = cells
                continue
            rec = dict(zip(hdr, cells))
            if not rec.get("SRC_ID", "").startswith("SRC_"):
                continue
            if rec.get("狀態") != "Active" or rec.get("來源等級") == "A":
                continue          # retired rows and Andy's own judgments excluded
            yield name, {k: rec.get(k, "") for k in KEEP}


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("terms", nargs="*")
    p.add_argument("--sheet")
    p.add_argument("--n", type=int, default=30)
    p.add_argument("--list-sheets", action="store_true")
    a = p.parse_args()
    path = ensure_clone()
    print(f"# source: {REPO} @ {os.path.basename(path)} (leads only; re-fetch originals)")
    if a.list_sheets:
        from collections import Counter
        for s, c in Counter(s for s, _ in rows(path, None)).items():
            print(s, c)
        return
    if not a.terms:
        sys.exit("give at least one search term")
    terms = [t.lower() for t in a.terms]
    n = 0
    for sheet, r in rows(path, a.sheet):
        blob = " ".join(r.values()).lower()
        if any(t in blob for t in terms):
            print(f"{sheet} | " + " | ".join(r[k] for k in KEEP))
            n += 1
            if n >= a.n:
                break
    print(f"# {n} lead(s)")


if __name__ == "__main__":
    main()
